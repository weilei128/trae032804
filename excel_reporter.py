import os
from datetime import datetime
from typing import Dict, Any, List
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import config


class ExcelReporter:
    def __init__(self, output_dir: str = None):
        self.output_dir = output_dir or config.OUTPUT_DIR
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir)
        
        self.header_font = Font(bold=True, color='FFFFFF')
        self.header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        self.header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        self.cell_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        self.center_alignment = Alignment(horizontal='center', vertical='center')
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        self.warning_fill = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')
        self.danger_fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
        self.success_fill = PatternFill(start_color='92D050', end_color='92D050', fill_type='solid')

    def apply_header_style(self, cell):
        cell.font = self.header_font
        cell.fill = self.header_fill
        cell.alignment = self.header_alignment
        cell.border = self.thin_border

    def apply_cell_style(self, cell, center=False):
        cell.alignment = self.center_alignment if center else self.cell_alignment
        cell.border = self.thin_border

    def auto_column_width(self, worksheet):
        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width

    def create_summary_sheet(self, workbook: Workbook, data: Dict[str, Any]):
        ws = workbook.active
        ws.title = "巡检概览"
        
        headers = ['服务器', 'IP地址', '巡检时间', '系统', '运行时间', 
                   'CPU负载(1/5/15min)', 'CPU使用率', '内存使用率', 
                   'Docker状态', '容器运行数']
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self.apply_header_style(cell)
        
        row = 2
        for server_name, server_data in data.items():
            sys_info = server_data.get('system', {})
            cpu_info = server_data.get('cpu', {})
            mem_info = server_data.get('memory', {})
            docker_info = server_data.get('docker_summary', {})
            inspection_time = server_data.get('inspection_time', '')
            
            cpu_load = f"{cpu_info.get('load_1min', 0):.2f}/{cpu_info.get('load_5min', 0):.2f}/{cpu_info.get('load_15min', 0):.2f}"
            
            ws.cell(row=row, column=1, value=server_name)
            ws.cell(row=row, column=2, value=sys_info.get('host', ''))
            ws.cell(row=row, column=3, value=inspection_time)
            ws.cell(row=row, column=4, value=sys_info.get('os', ''))
            ws.cell(row=row, column=5, value=sys_info.get('uptime', ''))
            ws.cell(row=row, column=6, value=cpu_load)
            
            cpu_usage_cell = ws.cell(row=row, column=7, value=f"{cpu_info.get('cpu_usage', 0):.1f}%")
            self.apply_cell_style(cpu_usage_cell, center=True)
            if cpu_info.get('cpu_usage', 0) > 80:
                cpu_usage_cell.fill = self.danger_fill
            elif cpu_info.get('cpu_usage', 0) > 60:
                cpu_usage_cell.fill = self.warning_fill
            
            mem_usage_cell = ws.cell(row=row, column=8, value=f"{mem_info.get('usage_percent', 0):.1f}%")
            self.apply_cell_style(mem_usage_cell, center=True)
            if mem_info.get('usage_percent', 0) > 80:
                mem_usage_cell.fill = self.danger_fill
            elif mem_info.get('usage_percent', 0) > 60:
                mem_usage_cell.fill = self.warning_fill
            
            docker_status = "运行中" if docker_info.get('running', False) else "未运行"
            ws.cell(row=row, column=9, value=docker_status)
            ws.cell(row=row, column=10, value=docker_info.get('containers_running', 0))
            
            for col in range(1, 11):
                self.apply_cell_style(ws.cell(row=row, column=col), center=True)
            
            row += 1
        
        self.auto_column_width(ws)

    def create_disk_sheet(self, workbook: Workbook, data: Dict[str, Any]):
        ws = workbook.create_sheet("磁盘信息")
        
        headers = ['服务器', '文件系统', '总容量', '已使用', '可用', '使用率', '挂载点']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self.apply_header_style(cell)
        
        row = 2
        for server_name, server_data in data.items():
            disk_info = server_data.get('disk', [])
            for disk in disk_info:
                ws.cell(row=row, column=1, value=server_name)
                ws.cell(row=row, column=2, value=disk.get('filesystem', ''))
                ws.cell(row=row, column=3, value=disk.get('size', ''))
                ws.cell(row=row, column=4, value=disk.get('used', ''))
                ws.cell(row=row, column=5, value=disk.get('available', ''))
                
                usage_cell = ws.cell(row=row, column=6, value=f"{disk.get('usage_percent', 0)}%")
                self.apply_cell_style(usage_cell, center=True)
                if disk.get('usage_percent', 0) > 80:
                    usage_cell.fill = self.danger_fill
                elif disk.get('usage_percent', 0) > 60:
                    usage_cell.fill = self.warning_fill
                
                ws.cell(row=row, column=7, value=disk.get('mount_point', ''))
                
                for col in range(1, 8):
                    self.apply_cell_style(ws.cell(row=row, column=col), center=True)
                
                row += 1
        
        self.auto_column_width(ws)

    def create_port_sheet(self, workbook: Workbook, data: Dict[str, Any]):
        ws = workbook.create_sheet("端口状态")
        
        headers = ['服务器', '端口', '状态', '服务程序']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self.apply_header_style(cell)
        
        row = 2
        for server_name, server_data in data.items():
            port_info = server_data.get('listening_ports', [])
            for port in port_info:
                ws.cell(row=row, column=1, value=server_name)
                ws.cell(row=row, column=2, value=port.get('port', ''))
                
                status_cell = ws.cell(row=row, column=3, value='LISTENING')
                self.apply_cell_style(status_cell, center=True)
                status_cell.fill = self.success_fill
                
                ws.cell(row=row, column=4, value=port.get('program', ''))
                
                for col in range(1, 5):
                    self.apply_cell_style(ws.cell(row=row, column=col), center=True)
                
                row += 1
        
        self.auto_column_width(ws)

    def create_process_sheet(self, workbook: Workbook, data: Dict[str, Any]):
        ws = workbook.create_sheet("进程状态")
        
        headers = ['服务器', '进程名', '状态', '实例数', 'PID列表']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self.apply_header_style(cell)
        
        row = 2
        for server_name, server_data in data.items():
            process_info = server_data.get('processes', [])
            for proc in process_info:
                ws.cell(row=row, column=1, value=server_name)
                ws.cell(row=row, column=2, value=proc.get('name', ''))
                
                status_cell = ws.cell(row=row, column=3, value=proc.get('status', ''))
                self.apply_cell_style(status_cell, center=True)
                if proc.get('is_running', False):
                    status_cell.fill = self.success_fill
                else:
                    status_cell.fill = self.warning_fill
                
                ws.cell(row=row, column=4, value=proc.get('count', 0))
                ws.cell(row=row, column=5, value=proc.get('pids', ''))
                
                for col in range(1, 6):
                    self.apply_cell_style(ws.cell(row=row, column=col), center=True)
                
                row += 1
        
        self.auto_column_width(ws)

    def create_docker_sheet(self, workbook: Workbook, data: Dict[str, Any]):
        ws = workbook.create_sheet("Docker容器")
        
        headers = ['服务器', '容器ID', '容器名', '镜像', '状态', '端口映射', '运行状态']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self.apply_header_style(cell)
        
        row = 2
        for server_name, server_data in data.items():
            containers = server_data.get('containers', [])
            for container in containers:
                ws.cell(row=row, column=1, value=server_name)
                ws.cell(row=row, column=2, value=container.get('id', ''))
                ws.cell(row=row, column=3, value=container.get('name', ''))
                ws.cell(row=row, column=4, value=container.get('image', ''))
                ws.cell(row=row, column=5, value=container.get('status', ''))
                ws.cell(row=row, column=6, value=container.get('ports', ''))
                
                state_cell = ws.cell(row=row, column=7, value=container.get('state', ''))
                self.apply_cell_style(state_cell, center=True)
                if container.get('is_running', False):
                    state_cell.fill = self.success_fill
                else:
                    state_cell.fill = self.warning_fill
                
                for col in range(1, 8):
                    self.apply_cell_style(ws.cell(row=row, column=col), center=True)
                
                row += 1
        
        self.auto_column_width(ws)

    def create_docker_images_sheet(self, workbook: Workbook, data: Dict[str, Any]):
        ws = workbook.create_sheet("Docker镜像")
        
        headers = ['服务器', '仓库', '标签', '镜像ID', '大小', '创建时间']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self.apply_header_style(cell)
        
        row = 2
        for server_name, server_data in data.items():
            images = server_data.get('images', [])
            for image in images:
                ws.cell(row=row, column=1, value=server_name)
                ws.cell(row=row, column=2, value=image.get('repository', ''))
                ws.cell(row=row, column=3, value=image.get('tag', ''))
                ws.cell(row=row, column=4, value=image.get('id', ''))
                ws.cell(row=row, column=5, value=image.get('size', ''))
                ws.cell(row=row, column=6, value=image.get('created', ''))
                
                for col in range(1, 7):
                    self.apply_cell_style(ws.cell(row=row, column=col), center=True)
                
                row += 1
        
        self.auto_column_width(ws)

    def create_history_sheet(self, workbook: Workbook, history_data: List[Dict[str, Any]]):
        ws = workbook.create_sheet("历史记录")
        
        headers = ['巡检时间', '服务器', 'CPU使用率', '内存使用率', '磁盘最高使用率', 
                   'Docker状态', '容器运行数', '异常项']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self.apply_header_style(cell)
        
        row = 2
        for record in history_data:
            ws.cell(row=row, column=1, value=record.get('inspection_time', ''))
            ws.cell(row=row, column=2, value=record.get('server_name', ''))
            ws.cell(row=row, column=3, value=f"{record.get('cpu_usage', 0):.1f}%")
            ws.cell(row=row, column=4, value=f"{record.get('memory_usage', 0):.1f}%")
            ws.cell(row=row, column=5, value=f"{record.get('disk_max_usage', 0)}%")
            ws.cell(row=row, column=6, value=record.get('docker_status', ''))
            ws.cell(row=row, column=7, value=record.get('containers_running', 0))
            ws.cell(row=row, column=8, value=record.get('abnormal_count', 0))
            
            for col in range(1, 9):
                self.apply_cell_style(ws.cell(row=row, column=col), center=True)
            
            row += 1
        
        self.auto_column_width(ws)

    def generate_report(self, data: Dict[str, Any], history_data: List[Dict[str, Any]] = None, 
                       filename: str = None) -> str:
        if filename is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"inspection_report_{timestamp}.xlsx"
        
        filepath = os.path.join(self.output_dir, filename)
        
        workbook = Workbook()
        
        self.create_summary_sheet(workbook, data)
        self.create_disk_sheet(workbook, data)
        self.create_port_sheet(workbook, data)
        self.create_process_sheet(workbook, data)
        self.create_docker_sheet(workbook, data)
        self.create_docker_images_sheet(workbook, data)
        
        if history_data:
            self.create_history_sheet(workbook, history_data)
        
        workbook.save(filepath)
        
        return filepath

    def update_report(self, filepath: str, data: Dict[str, Any], history_data: List[Dict[str, Any]]):
        workbook = Workbook()
        
        self.create_summary_sheet(workbook, data)
        self.create_disk_sheet(workbook, data)
        self.create_port_sheet(workbook, data)
        self.create_process_sheet(workbook, data)
        self.create_docker_sheet(workbook, data)
        self.create_docker_images_sheet(workbook, data)
        self.create_history_sheet(workbook, history_data)
        
        workbook.save(filepath)
        
        return filepath
